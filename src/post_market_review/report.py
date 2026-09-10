from __future__ import annotations

import html
from collections import defaultdict
from pathlib import Path
from typing import Any

from .research import inject_tokens
from .util import atomic_write_text


SECTION_TITLES = [
    ("S1", "一 当日市场定性"),
    ("S2", "二 盘面结构"),
    ("S3", "三 情绪观察"),
    ("S4", "四 晨报预测初步对照"),
    ("S5", "五 异动与消息面"),
    ("S6", "六 明日关注点"),
    ("S7", "七 数据源状态"),
]


def _fact_map(facts: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {item["evidence_id"]: item for item in facts}


def build_document(
    manifest: dict[str, Any],
    snapshot: dict[str, Any],
    facts: list[dict[str, Any]],
    groups: dict[str, list[str]],
    analysis: dict[str, Any],
    comparisons: list[dict[str, Any]],
    evidence: dict[str, Any],
    claims: list[dict[str, Any]],
    watch_items: list[dict[str, Any]],
    rejections: list[str],
) -> dict[str, Any]:
    fact_by_id = _fact_map(facts)
    claim_sections: dict[str, list[str]] = defaultdict(list)
    for claim in claims:
        claim_sections[claim["section_id"]].append(inject_tokens(claim["text_template"], evidence) + "〔证据：" + "、".join(claim["evidence_ids"]) + "〕")
    market = snapshot.get("market")
    sections: list[dict[str, Any]] = []
    for section_id, title in SECTION_TITLES:
        section: dict[str, Any] = {"section_id": section_id, "title": title, "paragraphs": [], "tables": []}
        if section_id == "S1":
            section["paragraphs"].append(analysis["market_label"])
            if market:
                rows = [[item["name"], item["close"] + "点", ("↑ +" if float(item["return_pct"]) > 0 else "↓ " if float(item["return_pct"]) < 0 else "") + item["return_pct"] + "%", item["open"], item["high"], item["low"]] for item in market["indexes"]]
                section["tables"].append({"headers": ["指数", "收盘", "涨跌幅", "开盘", "最高", "最低"], "rows": rows})
            else:
                section["paragraphs"].append("当日指数行情未获取。")
        elif section_id == "S2":
            section["paragraphs"].extend([
                fact_by_id[item]["label"] + "：" + fact_by_id[item]["display_value"] + "〔证据：" + item + "〕"
                for item in groups["structure"]
            ])
            if market:
                rows = [[str(index + 1), item["name"], ("↑ +" if float(item["return_pct"]) > 0 else "↓ ") + item["return_pct"] + "%", "当日强势"] for index, item in enumerate(market["industry_top5"])]
                rows += [[str(index + 1), item["name"], ("↑ +" if float(item["return_pct"]) > 0 else "↓ ") + item["return_pct"] + "%", "当日弱势"] for index, item in enumerate(market["industry_bottom5"])]
                section["tables"].append({"headers": ["排名", market["industry_classification"], "涨跌幅", "定位"], "rows": rows})
                if groups["anchors"]:
                    item = groups["anchors"][0]
                    section["paragraphs"].append(fact_by_id[item]["label"] + "：" + fact_by_id[item]["display_value"] + "。该列表只作为成交活跃锚点，不等同龙头。〔证据：" + item + "〕")
        elif section_id == "S3":
            section["paragraphs"].extend([
                fact_by_id[item]["label"] + "：" + fact_by_id[item]["display_value"] + "〔证据：" + item + "〕"
                for item in groups["emotion"]
            ])
            sentiment = analysis["sentiment"]
            if sentiment["status"] == "VALID":
                section["paragraphs"].append(f"情绪温度计：{sentiment['score']}，{sentiment['band']}（规则{sentiment['rule_version']}）。")
            else:
                section["paragraphs"].append("情绪温度计：未给出。炸板率或严格历史窗口缺失，不能以估算值替代。")
        elif section_id == "S4":
            rows = [[item["target_id"], item["forecast_display"], item["actual_display"], item["verdict"], "、".join(item["reason_codes"])] for item in comparisons]
            section["tables"].append({"headers": ["预测维度", "晨报预测", "实际走势", "初步判定", "说明"], "rows": rows})
            section["paragraphs"].append("本节仅做初步对照；正式评价由独立评价流程按预登记规则完成。")
        elif section_id == "S5":
            section["paragraphs"].append("当前未启用具备精确发布时间和外发许可的消息源，本节不以新闻摘要补写盘面原因。")
        elif section_id == "S6":
            if watch_items:
                for item in watch_items[:5]:
                    section["paragraphs"].append(inject_tokens(item["hypothesis_template"], evidence) + "〔证据：" + "、".join(item["evidence_ids"]) + f"；有效至{item['valid_until_session']}〕")
            elif market:
                top = groups["industries"][0] if groups["industries"] else None
                breadth = groups["structure"][1] if len(groups["structure"]) > 1 else None
                if top:
                    section["paragraphs"].append("观察当日强势行业能否延续相对强度并出现成交扩散。〔证据：" + top + "〕")
                if breadth:
                    section["paragraphs"].append("观察上涨家数与下跌家数是否改善，验证市场广度变化。〔证据：" + breadth + "〕")
            else:
                section["paragraphs"].append("缺少合格市场事实，本次不生成次日观察项。")
        elif section_id == "S7":
            rows = [[item["provider"], item["status"], item["reason"]] for item in snapshot["source_states"]]
            rows.append(["DeepSeek", manifest["narrative_status"], manifest.get("llm_error", "")])
            section["tables"].append({"headers": ["数据源或组件", "状态", "说明"], "rows": rows})
            section["paragraphs"].append("数据来源于 Wind Alice 万得金融数据服务。自动化研究产出，不构成投资建议。")
        section["paragraphs"].extend(claim_sections.get(section_id, []))
        sections.append(section)
    limitations = []
    if snapshot["cutoff_status"] != "ON_TIME":
        limitations.append("本次数据取得时间晚于19:45冻结点，报告保留真实可得时间并标记迟到采集。")
    if rejections:
        limitations.append(f"DeepSeek输出中有{len(rejections)}条未通过结构或证据校验，未进入正文。")
    limitations.append("行业榜为Wind行业板块口径；未验证为申万一级行业。")
    return {
        "schema_version": "review.document.v1",
        "title": f"{int(manifest['report_date'][0:4])}年{int(manifest['report_date'][5:7])}月{int(manifest['report_date'][8:10])}日A股每日收盘复盘",
        "report_date": manifest["report_date"],
        "status": manifest["report_status"],
        "narrative_status": manifest["narrative_status"],
        "cutoff_at": snapshot["cutoff_at"],
        "created_at": manifest["created_at"],
        "sections": sections,
        "source_notes": ["Wind为收盘结构化事实主源；指数复核属于同供应商一致性检查，不构成独立双源。"],
        "limitations": limitations,
    }


def to_markdown(document: dict[str, Any]) -> str:
    lines = [
        "# " + document["title"],
        "",
        f"> 报告状态：{document['status']}  ",
        f"> 研究状态：{document['narrative_status']}  ",
        f"> 证据冻结点：{document['cutoff_at']}  ",
        "> 自动化研究产出，不构成投资建议",
    ]
    for section in document["sections"]:
        lines += ["", "## " + section["title"], ""]
        for paragraph in section["paragraphs"]:
            lines += [paragraph, ""]
        for table in section["tables"]:
            lines.append("| " + " | ".join(table["headers"]) + " |")
            lines.append("|" + "|".join("---" for _ in table["headers"]) + "|")
            for row in table["rows"]:
                lines.append("| " + " | ".join(str(cell).replace("|", "\\|") for cell in row) + " |")
            lines.append("")
    lines += ["## 限制说明", ""]
    lines += ["- " + item for item in document["limitations"]]
    return "\n".join(lines).rstrip() + "\n"


def to_html(document: dict[str, Any]) -> str:
    parts = [
        "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width,initial-scale=1'>",
        f"<title>{html.escape(document['title'])}</title>",
        "<style>body{font-family:'Microsoft YaHei','Noto Sans CJK SC',sans-serif;color:#171717;max-width:1180px;margin:40px auto;padding:0 28px;line-height:1.65}"
        "h1,h2{color:#000}h1{font-size:30px;margin-bottom:8px}h2{font-size:21px;margin-top:34px}"
        ".meta{color:#555}.up{color:#C62828;font-weight:600}.down{color:#15803D;font-weight:600}"
        "table{border-collapse:collapse;width:100%;margin:14px 0 24px;font-size:14px}th{background:#163A5F;color:white}"
        "th,td{border:1px solid #D9D9D9;padding:9px 11px;vertical-align:middle}tr:nth-child(even) td{background:#F4F7FA}"
        "ul{padding-left:22px}</style></head><body>",
        f"<h1>{html.escape(document['title'])}</h1>",
        f"<p class='meta'>报告状态：{html.escape(document['status'])}　研究状态：{html.escape(document['narrative_status'])}<br>"
        f"证据冻结点：{html.escape(document['cutoff_at'])}<br>自动化研究产出，不构成投资建议</p>",
    ]
    for section in document["sections"]:
        parts.append(f"<h2>{html.escape(section['title'])}</h2>")
        for paragraph in section["paragraphs"]:
            css = "up" if "↑ +" in paragraph else "down" if "↓ -" in paragraph else ""
            parts.append(f"<p class='{css}'>{html.escape(paragraph)}</p>")
        for table in section["tables"]:
            parts.append("<table><thead><tr>" + "".join(f"<th>{html.escape(str(cell))}</th>" for cell in table["headers"]) + "</tr></thead><tbody>")
            for row in table["rows"]:
                cells = []
                for cell in row:
                    value = str(cell)
                    css = "up" if "↑ +" in value else "down" if "↓ -" in value else ""
                    cells.append(f"<td class='{css}'>{html.escape(value)}</td>")
                parts.append("<tr>" + "".join(cells) + "</tr>")
            parts.append("</tbody></table>")
    parts.append("<h2>限制说明</h2><ul>" + "".join(f"<li>{html.escape(item)}</li>" for item in document["limitations"]) + "</ul>")
    parts.append("</body></html>")
    return "".join(parts)


def _set_cell_borders(cell: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)


def to_docx(document: dict[str, Any], path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    for style_name, size, bold in (("Normal", 10.5, False), ("Title", 22, True), ("Heading 1", 15, True)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(document["title"])
    intro = doc.add_paragraph()
    intro.add_run(f"报告状态：{document['status']}　研究状态：{document['narrative_status']}\n").bold = True
    intro.add_run(f"证据冻结点：{document['cutoff_at']}\n自动化研究产出，不构成投资建议")
    for item in document["sections"]:
        heading = doc.add_paragraph(item["title"], style="Heading 1")
        heading.paragraph_format.keep_with_next = True
        for paragraph in item["paragraphs"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(paragraph)
            if "↑ +" in paragraph:
                run.font.color.rgb = RGBColor(198, 40, 40)
            elif "↓ -" in paragraph:
                run.font.color.rgb = RGBColor(21, 128, 61)
        for table_data in item["tables"]:
            table = doc.add_table(rows=1, cols=len(table_data["headers"]))
            table.autofit = True
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for index, header in enumerate(table_data["headers"]):
                cell = table.rows[0].cells[index]
                cell.text = str(header)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "163A5F")
                cell._tc.get_or_add_tcPr().append(shading)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                    run.font.size = Pt(9)
            for row_index, values in enumerate(table_data["rows"]):
                cells = table.add_row().cells
                for index, value in enumerate(values):
                    cell = cells[index]
                    cell.text = str(value)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if row_index % 2:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "F4F7FA")
                        cell._tc.get_or_add_tcPr().append(shading)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
                        run.font.name = "Microsoft YaHei"
                        run_properties = run._element.get_or_add_rPr()
                        run_fonts = run_properties.find(qn("w:rFonts"))
                        if run_fonts is None:
                            run_fonts = OxmlElement("w:rFonts")
                            run_properties.append(run_fonts)
                        run_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                        if "↑ +" in str(value):
                            run.font.color.rgb = RGBColor(198, 40, 40)
                        elif "↓ -" in str(value):
                            run.font.color.rgb = RGBColor(21, 128, 61)
            for row in table.rows:
                for cell in row.cells:
                    _set_cell_borders(cell)
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
    doc.add_paragraph("限制说明", style="Heading 1")
    for limitation in document["limitations"]:
        doc.add_paragraph(limitation, style="List Bullet")
    doc.core_properties.title = document["title"]
    doc.core_properties.subject = "A股每日收盘复盘"
    doc.core_properties.author = "Richard Daily Done"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def render_all(document: dict[str, Any], directory: Path, basename: str) -> dict[str, Path]:
    paths = {
        "markdown": directory / f"{basename}.md",
        "html": directory / f"{basename}.html",
        "docx": directory / f"{basename}.docx",
    }
    atomic_write_text(paths["markdown"], to_markdown(document))
    atomic_write_text(paths["html"], to_html(document))
    to_docx(document, paths["docx"])
    return paths
